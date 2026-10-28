"""Stage 1 & Stage 7: Multi-Source Data Ingestion Service

Handles ingestion from diverse data sources into standardized pandas DataFrame(s):
1. CSV / Excel files (.csv, .xlsx, .xls)
2. Live Database connections (PostgreSQL, MySQL, SQLite via SQLAlchemy) with strict SELECT-only SQL validation
3. PDF documents (.pdf) with multi-table extraction and unstructured text fallback
4. PowerPoint presentations (.pptx) with slide table extraction and per-slide text notes

Normalizes every source into the unified DatasetMetadata contract and session store.
"""

import io
import math
import re
import uuid
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional, Tuple
import numpy as np
import pandas as pd
from sqlalchemy import create_engine, text

import pdfplumber
from pptx import Presentation

from backend.app.core.config import settings
from backend.app.models.schemas import DatasetMetadata, ColumnInfo


class IngestionError(Exception):
    """Base exception for data ingestion errors."""
    pass


class InvalidFileTypeError(IngestionError):
    """Raised when an unsupported file extension is uploaded."""
    pass


class EmptyFileError(IngestionError):
    """Raised when an uploaded file is empty or contains no data."""
    pass


class CorruptedFileError(IngestionError):
    """Raised when a file cannot be parsed or decoded."""
    pass


class SQLValidationError(IngestionError):
    """Raised when a supplied SQL query violates security constraints (e.g. non-SELECT or destructive)."""
    pass


class DatabaseConnectionError(IngestionError):
    """Raised when connection to an external database fails."""
    pass


# Strict list of disallowed SQL operational keywords
FORBIDDEN_SQL_KEYWORDS = {
    "INSERT", "UPDATE", "DELETE", "DROP", "ALTER", "TRUNCATE",
    "EXEC", "EXECUTE", "CREATE", "REPLACE", "GRANT", "REVOKE",
    "MERGE", "UPSERT", "CALL", "COPY", "VACUUM", "ATTACH", "DETACH"
}


def validate_select_query(query: str) -> str:
    """Validate that a user query is strictly SELECT-only and safe to execute.

    Args:
        query: Raw SQL query string.

    Returns:
        Cleaned SQL query string.

    Raises:
        SQLValidationError: If the query is empty, chained, non-SELECT, or contains forbidden keywords.
    """
    cleaned = (query or "").strip()
    if not cleaned:
        raise SQLValidationError("SQL query cannot be empty.")

    # Remove SQL single-line and multi-line comments
    no_comments = re.sub(r"--[^\r\n]*", " ", cleaned)
    no_comments = re.sub(r"/\*.*?\*/", " ", no_comments, flags=re.DOTALL).strip()

    # Reject statement chaining with semicolons (allowing optional trailing semicolon)
    without_end_semi = no_comments.rstrip(";").strip()
    if ";" in without_end_semi:
        raise SQLValidationError("Multiple SQL statements or chained semicolons are prohibited.")

    # Must start with SELECT or WITH (for CTE expressions)
    if not re.match(r"^\s*(SELECT|WITH)\b", no_comments, re.IGNORECASE):
        raise SQLValidationError("Security violation: Only SELECT queries (or WITH ... SELECT CTEs) are permitted.")

    # Check for forbidden keywords as distinct word tokens
    tokens = re.findall(r"\b[A-Za-z_][A-Za-z0-9_]*\b", no_comments)
    for token in tokens:
        if token.upper() in FORBIDDEN_SQL_KEYWORDS:
            raise SQLValidationError(
                f"Forbidden SQL operation detected: '{token.upper()}'. Only SELECT queries are permitted."
            )

    return cleaned


def validate_table_name(table_name: str) -> str:
    """Validate that table_name is a valid SQL identifier (alphanumeric and underscores only).

    Args:
        table_name: Target table name string.

    Returns:
        Validated table name string.

    Raises:
        SQLValidationError: If table identifier is empty or contains invalid characters.
    """
    cleaned = (table_name or "").strip()
    if not cleaned:
        raise SQLValidationError("Table name cannot be empty.")

    # Support optional schema.table syntax (e.g., 'public.sales')
    parts = cleaned.split(".")
    for part in parts:
        if not re.match(r"^[a-zA-Z_][a-zA-Z0-9_]*$", part):
            raise SQLValidationError(
                f"Invalid table identifier '{table_name}'. Table names must be alphanumeric identifiers."
            )
    return cleaned


class DatasetSessionStore:
    """In-memory session store caching loaded DataFrames, additional tables, and metadata."""

    def __init__(self):
        self._dataframes: Dict[str, pd.DataFrame] = {}
        self._metadata: Dict[str, DatasetMetadata] = {}
        self._extra_tables: Dict[str, List[pd.DataFrame]] = {}
        self._latest_id: Optional[str] = None

    def save(
        self,
        file_id: str,
        df: pd.DataFrame,
        metadata: DatasetMetadata,
        session_id: Optional[str] = None,
        extra_tables: Optional[List[pd.DataFrame]] = None,
    ) -> None:
        """Store the DataFrame and associated metadata."""
        self._dataframes[file_id] = df
        self._metadata[file_id] = metadata
        self._latest_id = file_id
        if extra_tables:
            self._extra_tables[file_id] = extra_tables

        if session_id and session_id != file_id:
            self._dataframes[session_id] = df
            self._metadata[session_id] = metadata
            if extra_tables:
                self._extra_tables[session_id] = extra_tables

    def get_dataframe(self, file_id: Optional[str] = None) -> Optional[pd.DataFrame]:
        """Retrieve stored DataFrame by file_id, or return the latest loaded dataset if omitted."""
        if not file_id or file_id in ("latest", "current", "default"):
            file_id = self._latest_id
        if not file_id:
            return None
        return self._dataframes.get(file_id)

    def get_metadata(self, file_id: Optional[str] = None) -> Optional[DatasetMetadata]:
        """Retrieve stored DatasetMetadata by file_id, or return latest metadata if omitted."""
        if not file_id or file_id in ("latest", "current", "default"):
            file_id = self._latest_id
        if not file_id:
            return None
        return self._metadata.get(file_id)

    def get_extra_tables(self, file_id: Optional[str] = None) -> List[pd.DataFrame]:
        """Retrieve additional tables if source contained multiple tables."""
        if not file_id or file_id in ("latest", "current", "default"):
            file_id = self._latest_id
        if not file_id:
            return []
        return self._extra_tables.get(file_id, [])

    def get_latest_id(self) -> Optional[str]:
        """Retrieve identifier of the most recently ingested dataset."""
        return self._latest_id

    def get_latest_dataframe(self) -> Optional[pd.DataFrame]:
        """Retrieve the currently/most recently loaded DataFrame."""
        return self.get_dataframe(self._latest_id)

    def get_latest_metadata(self) -> Optional[DatasetMetadata]:
        """Retrieve metadata for the currently/most recently loaded dataset."""
        return self.get_metadata(self._latest_id)

    def has(self, file_id: str) -> bool:
        """Check if dataset exists in store."""
        if not file_id or file_id in ("latest", "current", "default"):
            return self._latest_id is not None and self._latest_id in self._dataframes
        return file_id in self._dataframes

    def list_all(self) -> List[str]:
        """List all active dataset file_ids."""
        return list(self._dataframes.keys())

    def clear(self) -> None:
        """Clear all stored datasets."""
        self._dataframes.clear()
        self._metadata.clear()
        self._extra_tables.clear()
        self._latest_id = None


# Global in-memory dataset store
dataset_store = DatasetSessionStore()


def _sanitize_for_json(val: Any) -> Any:
    """Convert numpy/pandas types to standard JSON-serializable Python types."""
    if val is None or pd.isna(val):
        return None
    if isinstance(val, (np.integer, int)):
        return int(val)
    if isinstance(val, (np.floating, float)):
        if math.isnan(val) or math.isinf(val):
            return None
        return float(val)
    if isinstance(val, (np.bool_, bool)):
        return bool(val)
    if isinstance(val, (pd.Timestamp, np.datetime64)):
        return str(val)
    return str(val)


class DataIngestionService:
    """Service handling multi-source ingestion (CSV, Excel, Database, PDF, PPTX) and schema extraction."""

    ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls", ".pdf", ".pptx"}

    def __init__(self, store: Optional[DatasetSessionStore] = None, upload_dir: Optional[Path] = None):
        self.store = store or dataset_store
        self.upload_dir = upload_dir or settings.UPLOAD_DIR

    def ingest_bytes(
        self,
        filename: str,
        content: bytes,
        file_id: Optional[str] = None,
        session_id: Optional[str] = None,
    ) -> DatasetMetadata:
        """Parse raw uploaded file bytes, extract schema metadata, and cache in session store.

        Supports: .csv, .xlsx, .xls, .pdf, .pptx

        Args:
            filename: Original uploaded file name (e.g. 'sales.csv', 'report.pdf')
            content: Raw file content bytes
            file_id: Optional custom identifier.
            session_id: Optional session identifier alias.

        Returns:
            DatasetMetadata normalized across all formats.
        """
        ext = Path(filename).suffix.lower()
        if ext not in self.ALLOWED_EXTENSIONS:
            raise InvalidFileTypeError(
                f"Unsupported file format '{ext}'. Supported formats: CSV (.csv), Excel (.xlsx, .xls), PDF (.pdf), PowerPoint (.pptx)."
            )

        if not content or len(content.strip()) == 0:
            raise EmptyFileError("The uploaded file is empty (0 bytes).")

        assigned_id = file_id or session_id

        # 1. Tabular Files: CSV and Excel
        if ext in (".csv", ".xlsx", ".xls"):
            source_type: Literal["csv", "excel"] = "csv" if ext == ".csv" else "excel"
            df = self._read_into_dataframe(content, ext)

            if df.empty or len(df.columns) == 0:
                raise EmptyFileError("The uploaded file contains no data rows or columns.")

            df.columns = [str(col).strip() for col in df.columns]
            metadata = self._profile_dataframe(
                df=df,
                filename=filename,
                file_id=assigned_id,
                source_type=source_type,
                has_structured_tables=True,
                has_unstructured_text=False,
            )
            self.store.save(metadata.file_id, df, metadata, session_id=session_id)

        # 2. PDF Documents
        elif ext == ".pdf":
            tables, raw_text = self._read_pdf(content, filename)
            has_tables = len(tables) > 0

            if has_tables:
                primary_df = tables[0]
                table_summaries = [
                    {"table_index": idx, "row_count": len(t), "columns": list(t.columns)}
                    for idx, t in enumerate(tables)
                ]
                metadata = self._profile_dataframe(
                    df=primary_df,
                    filename=filename,
                    file_id=assigned_id,
                    source_type="pdf",
                    has_structured_tables=True,
                    has_unstructured_text=bool(raw_text),
                    unstructured_text=raw_text,
                    tables=table_summaries,
                    table_count=len(tables),
                )
                self.store.save(
                    metadata.file_id,
                    primary_df,
                    metadata,
                    session_id=session_id,
                    extra_tables=tables[1:] if len(tables) > 1 else None,
                )
            else:
                # PDF has NO extractable tables -> unstructured text only
                empty_df = pd.DataFrame()
                metadata = self._profile_dataframe(
                    df=empty_df,
                    filename=filename,
                    file_id=assigned_id,
                    source_type="pdf",
                    has_structured_tables=False,
                    has_unstructured_text=True,
                    unstructured_text=raw_text or "No text or tables found in PDF document.",
                    tables=[],
                    table_count=0,
                    warning="No extractable tables found in PDF. Extracted unstructured text only; table operations cannot be executed.",
                )
                self.store.save(metadata.file_id, empty_df, metadata, session_id=session_id)

        # 3. PowerPoint Presentations
        elif ext == ".pptx":
            tables, slide_notes, raw_text = self._read_pptx(content, filename)
            has_tables = len(tables) > 0

            if has_tables:
                primary_df = tables[0]
                table_summaries = [
                    {"table_index": idx, "row_count": len(t), "columns": list(t.columns)}
                    for idx, t in enumerate(tables)
                ]
                metadata = self._profile_dataframe(
                    df=primary_df,
                    filename=filename,
                    file_id=assigned_id,
                    source_type="pptx",
                    has_structured_tables=True,
                    has_unstructured_text=bool(raw_text),
                    unstructured_text=raw_text,
                    unstructured_notes=slide_notes,
                    tables=table_summaries,
                    table_count=len(tables),
                )
                self.store.save(
                    metadata.file_id,
                    primary_df,
                    metadata,
                    session_id=session_id,
                    extra_tables=tables[1:] if len(tables) > 1 else None,
                )
            else:
                # PPTX has NO tables -> unstructured slide text only
                empty_df = pd.DataFrame()
                metadata = self._profile_dataframe(
                    df=empty_df,
                    filename=filename,
                    file_id=assigned_id,
                    source_type="pptx",
                    has_structured_tables=False,
                    has_unstructured_text=True,
                    unstructured_text=raw_text or "No text or tables found in presentation.",
                    unstructured_notes=slide_notes,
                    tables=[],
                    table_count=0,
                    warning="No extractable tables found in presentation slides. Extracted unstructured text only; table operations cannot be executed.",
                )
                self.store.save(metadata.file_id, empty_df, metadata, session_id=session_id)

        # Persist uploaded bytes to disk
        try:
            saved_filename = f"{metadata.file_id}_{filename}"
            disk_path = self.upload_dir / saved_filename
            with open(disk_path, "wb") as f:
                f.write(content)
        except Exception:
            pass

        return metadata

    def ingest_database(
        self,
        connection_string: str,
        table_name: Optional[str] = None,
        query: Optional[str] = None,
        file_id: Optional[str] = None,
        session_id: Optional[str] = None,
    ) -> DatasetMetadata:
        """Connect to live database (PostgreSQL, MySQL, SQLite) and ingest into DataFrame.

        Enforces strict SQL validation (SELECT-only, rejects DDL/DML/destructive operations).

        Args:
            connection_string: SQLAlchemy URI.
            table_name: Target table name to ingest.
            query: Custom SELECT-only SQL query to execute.
            file_id: Optional custom identifier.
            session_id: Optional session identifier alias.

        Returns:
            DatasetMetadata normalized with source_type='db'.

        Raises:
            SQLValidationError: If the query fails safety checks.
            DatabaseConnectionError: If database connection or execution fails.
        """
        if not table_name and not query:
            raise SQLValidationError("Either 'table_name' or 'query' must be provided for database ingestion.")

        if query:
            validated_query = validate_select_query(query)
            display_name = f"db:query:{validated_query[:25]}..."
        else:
            safe_table = validate_table_name(table_name)
            validated_query = f"SELECT * FROM {safe_table}"
            display_name = f"db:table:{safe_table}"

        try:
            engine = create_engine(connection_string)
            with engine.connect() as conn:
                df = pd.read_sql_query(text(validated_query), con=conn)
        except SQLValidationError:
            raise
        except Exception as e:
            raise DatabaseConnectionError(f"Database query execution failed: {str(e)}")

        # Clean column names
        df.columns = [str(col).strip() for col in df.columns]

        assigned_id = file_id or session_id
        metadata = self._profile_dataframe(
            df=df,
            filename=display_name,
            file_id=assigned_id,
            source_type="db",
            has_structured_tables=True,
            has_unstructured_text=False,
            tables=[{"table_name": table_name or "custom_query", "row_count": len(df), "columns": list(df.columns)}],
            table_count=1,
        )

        self.store.save(metadata.file_id, df, metadata, session_id=session_id)
        return metadata

    def _read_into_dataframe(self, content: bytes, ext: str) -> pd.DataFrame:
        """Read bytes into DataFrame based on CSV/Excel extension with fallback encodings."""
        if ext == ".csv":
            for encoding in ("utf-8", "latin1", "cp1252"):
                try:
                    return pd.read_csv(io.BytesIO(content), encoding=encoding)
                except UnicodeDecodeError:
                    continue
                except pd.errors.EmptyDataError:
                    raise EmptyFileError("The uploaded CSV file contains no data or columns.")
                except Exception as e:
                    raise CorruptedFileError(f"Failed to parse CSV file: {str(e)}")
            raise CorruptedFileError("Failed to decode CSV file with supported encodings (UTF-8, Latin-1).")

        elif ext in (".xlsx", ".xls"):
            try:
                return pd.read_excel(io.BytesIO(content))
            except Exception as e:
                raise CorruptedFileError(f"Failed to parse Excel file: {str(e)}")

        raise InvalidFileTypeError(f"Unsupported file type: {ext}")

    def _read_pdf(self, content: bytes, filename: str) -> Tuple[List[pd.DataFrame], Optional[str]]:
        """Extract tables and raw text from PDF content bytes."""
        extracted_tables: List[pd.DataFrame] = []
        page_texts: List[str] = []

        try:
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for p_idx, page in enumerate(pdf.pages, 1):
                    # 1. Extract text
                    text_str = page.extract_text()
                    if text_str and text_str.strip():
                        page_texts.append(f"--- Page {p_idx} ---\n{text_str.strip()}")

                    # 2. Extract tables
                    tables = page.extract_tables()
                    for raw_table in tables:
                        if not raw_table:
                            continue
                        # Remove empty rows
                        clean_rows = [[cell.strip() if cell else "" for cell in row] for row in raw_table if any(row)]
                        if len(clean_rows) >= 2:
                            header = clean_rows[0]
                            data = clean_rows[1:]
                            # Ensure unique column names
                            cols = []
                            for idx, h in enumerate(header):
                                col_name = h if h else f"col_{idx}"
                                if col_name in cols:
                                    col_name = f"{col_name}_{idx}"
                                cols.append(col_name)
                            table_df = pd.DataFrame(data, columns=cols)
                            table_df.columns = [str(c).strip() for c in table_df.columns]
                            extracted_tables.append(table_df)
                        elif len(clean_rows) == 1:
                            cols = [f"col_{i}" for i in range(len(clean_rows[0]))]
                            table_df = pd.DataFrame(clean_rows, columns=cols)
                            table_df.columns = [str(c).strip() for c in table_df.columns]
                            extracted_tables.append(table_df)
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse PDF document: {str(e)}")

        full_unstructured_text = "\n\n".join(page_texts).strip() if page_texts else None
        return extracted_tables, full_unstructured_text

    def _read_pptx(
        self,
        content: bytes,
        filename: str
    ) -> Tuple[List[pd.DataFrame], List[Dict[str, Any]], Optional[str]]:
        """Extract embedded tables and slide text from PPTX bytes."""
        extracted_tables: List[pd.DataFrame] = []
        slide_notes: List[Dict[str, Any]] = []
        slide_texts_combined: List[str] = []

        try:
            prs = Presentation(io.BytesIO(content))
            for slide_idx, slide in enumerate(prs.slides, 1):
                slide_shape_texts: List[str] = []
                for shape in slide.shapes:
                    if shape.has_text_frame and shape.text_frame:
                        t = shape.text_frame.text.strip()
                        if t:
                            slide_shape_texts.append(t)
                    if shape.has_table and shape.table:
                        table = shape.table
                        grid = []
                        for row in table.rows:
                            row_cells = [cell.text.strip() for cell in row.cells]
                            if any(row_cells):
                                grid.append(row_cells)
                        if len(grid) >= 2:
                            header = grid[0]
                            data = grid[1:]
                            cols = []
                            for idx, h in enumerate(header):
                                col_name = h if h else f"col_{idx}"
                                if col_name in cols:
                                    col_name = f"{col_name}_{idx}"
                                cols.append(col_name)
                            table_df = pd.DataFrame(data, columns=cols)
                            table_df.columns = [str(c).strip() for c in table_df.columns]
                            extracted_tables.append(table_df)
                        elif len(grid) == 1:
                            cols = [f"col_{i}" for i in range(len(grid[0]))]
                            table_df = pd.DataFrame(grid, columns=cols)
                            table_df.columns = [str(c).strip() for c in table_df.columns]
                            extracted_tables.append(table_df)

                slide_text_content = "\n".join(slide_shape_texts).strip()
                if slide_text_content:
                    slide_notes.append({"slide": slide_idx, "text": slide_text_content})
                    slide_texts_combined.append(f"--- Slide {slide_idx} ---\n{slide_text_content}")
        except Exception as e:
            raise CorruptedFileError(f"Failed to parse PowerPoint presentation: {str(e)}")

        full_unstructured_text = "\n\n".join(slide_texts_combined).strip() if slide_texts_combined else None
        return extracted_tables, slide_notes, full_unstructured_text

    def _profile_dataframe(
        self,
        df: pd.DataFrame,
        filename: str,
        file_id: Optional[str] = None,
        source_type: Literal["csv", "excel", "db", "pdf", "pptx"] = "csv",
        has_structured_tables: bool = True,
        has_unstructured_text: bool = False,
        unstructured_text: Optional[str] = None,
        unstructured_notes: Optional[List[Dict[str, Any]]] = None,
        tables: Optional[List[Dict[str, Any]]] = None,
        table_count: int = 1,
        warning: Optional[str] = None,
    ) -> DatasetMetadata:
        """Compute column statistics, data types, nulls, and preview rows across all sources."""
        row_count = int(len(df))
        col_count = int(len(df.columns))
        assigned_id = file_id or uuid.uuid4().hex[:12]

        columns_profile: List[ColumnInfo] = []
        missing_dict: Dict[str, int] = {}

        if not df.empty and col_count > 0:
            for col_name in df.columns:
                series = df[col_name]
                missing_cnt = int(series.isna().sum())
                missing_dict[str(col_name)] = missing_cnt
                missing_pct = round((missing_cnt / row_count) * 100, 2) if row_count > 0 else 0.0

                # Gather sample values (up to 3 unique non-null values)
                valid_samples = series.dropna().unique()[:3]
                sample_vals = [_sanitize_for_json(v) for v in valid_samples]

                columns_profile.append(
                    ColumnInfo(
                        name=str(col_name),
                        dtype=str(series.dtype),
                        missing_count=missing_cnt,
                        missing_percentage=missing_pct,
                        sample_values=sample_vals,
                    )
                )

        # Preview rows (first 5)
        preview_rows: List[Dict[str, Any]] = []
        if not df.empty and col_count > 0:
            preview_slice = df.head(5).copy()
            for record in preview_slice.to_dict(orient="records"):
                clean_record = {k: _sanitize_for_json(v) for k, v in record.items()}
                preview_rows.append(clean_record)

        return DatasetMetadata(
            file_id=assigned_id,
            filename=filename,
            row_count=row_count,
            column_count=col_count,
            columns=columns_profile,
            preview_rows=preview_rows,
            missing_values=missing_dict,
            source_type=source_type,
            has_structured_tables=has_structured_tables,
            has_unstructured_text=has_unstructured_text,
            unstructured_text=unstructured_text,
            unstructured_notes=unstructured_notes or [],
            tables=tables or [],
            table_count=table_count,
            warning=warning,
        )

    def load_dataset(self, file_id: Optional[str] = None) -> pd.DataFrame:
        """Load stored dataset for a given file_id or currently loaded dataset if omitted."""
        df = self.store.get_dataframe(file_id)
        if df is None:
            id_repr = f"'{file_id}'" if file_id else "(latest/current)"
            raise KeyError(f"Dataset with ID {id_repr} not found in active session store.")
        return df

    def get_metadata(self, file_id: Optional[str] = None) -> Optional[DatasetMetadata]:
        """Retrieve dataset metadata for a given file_id or currently loaded dataset if omitted."""
        return self.store.get_metadata(file_id)

    def get_current_dataset(self) -> pd.DataFrame:
        """Retrieve the currently loaded dataset without needing an explicit ID."""
        return self.load_dataset(None)

    def get_current_metadata(self) -> Optional[DatasetMetadata]:
        """Retrieve the metadata for the currently loaded dataset."""
        return self.store.get_latest_metadata()


# Global service instance
ingestion_service = DataIngestionService()
